# -*- coding: utf-8 -*-
"""check_pattern_fields.py — 校验 profile「已验证模式」声明的字段是否仍存在于 config_root。

用法:
    python check_pattern_fields.py [--config-root PATH]

读取同目录 ../profiles/timemachine.md 中每个模式 fields: 块声明的
workbook/sheet/field_or_cell，逐表读取第 2 行（英文字段行）并核对字段存在性。

输出: 每个模式一行 pass/fail；fail 列出缺失字段。
返回码: 0=全部通过; 1=有字段缺失; 2=profile 或 config_root 不可用。
"""
import argparse
import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
PROFILE = HERE.parent / "profiles" / "timemachine.md"


def parse_patterns(profile_text: str):
    """从 profile 提取 (pattern_name, [(workbook, sheet, [fields])...])。

    fields 块格式:
        fields:
          - source_root: config_root
            workbook: A009-领袖表.xlsx
            sheet: 英雄基础表
            field_or_cell: quality,type
    """
    patterns = []
    cur_name, cur_sources = None, []
    in_fields = False
    src = {}
    for line in profile_text.splitlines():
        m = re.match(r"^### `([^`]+)`", line)
        if m:
            if cur_name and cur_sources:
                patterns.append((cur_name, cur_sources))
            cur_name, cur_sources, src, in_fields = m.group(1), [], {}, False
            continue
        if line.strip() == "fields:":
            in_fields = True
            continue
        if in_fields and line.startswith("```"):
            if src:
                cur_sources.append(src)
            src, in_fields = {}, False
            continue
        if in_fields:
            m = re.match(r"^\s+- source_root:", line)
            if m and src:
                cur_sources.append(src)
                src = {}
                continue
            m = re.match(r"^\s+workbook:\s*(.+?)\s*$", line)
            if m:
                src["workbook"] = m.group(1)
                continue
            m = re.match(r"^\s+sheet:\s*(.+?)\s*$", line)
            if m:
                src["sheet"] = m.group(1)
                continue
            m = re.match(r"^\s+field_or_cell:\s*(.+?)\s*$", line)
            if m:
                src["fields"] = [f.strip() for f in m.group(1).split(",") if f.strip()]
                continue
    if cur_name and cur_sources:
        patterns.append((cur_name, cur_sources))
    return patterns


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--config-root", default=None)
    args = ap.parse_args()

    if not PROFILE.is_file():
        print(f"profile 不存在: {PROFILE}", file=sys.stderr)
        return 2
    root = args.config_root
    if not root:
        local = PROFILE.parent / "timemachine.local.yaml"
        m = re.search(r"config_root:\s*'([^']+)'", local.read_text(encoding="utf-8")) if local.is_file() else None
        root = m.group(1) if m else None
    if not root or not os.path.isdir(root):
        print(f"config_root 不可用: {root}", file=sys.stderr)
        return 2

    patterns = parse_patterns(PROFILE.read_text(encoding="utf-8"))
    if not patterns:
        print("profile 中未解析到模式 fields 块", file=sys.stderr)
        return 2

    # 只读用到的 workbook，按需加载
    wb_cache = {}
    fails = 0
    for name, sources in patterns:
        missing = []
        for src in sources:
            wb_name = src.get("workbook")
            if wb_name is None or "（" in wb_name:
                continue  # 无实体来源的约定章节
            path = os.path.join(root, wb_name)
            if not os.path.isfile(path):
                missing.append(f"{wb_name}: 文件缺失")
                continue
            if wb_name not in wb_cache:
                try:
                    wb_cache[wb_name] = load_workbook(path, read_only=True, data_only=False)
                except Exception as e:
                    missing.append(f"{wb_name}: 读取失败 {e}")
                    continue
            wb = wb_cache[wb_name]
            sheet = src.get("sheet")
            if sheet not in wb.sheetnames:
                missing.append(f"{wb_name}!{sheet}: sheet 缺失")
                continue
            ws = wb[sheet]
            header = {ws.cell(2, c).value for c in range(1, ws.max_column + 1)}
            for f in src.get("fields", []):
                if f not in header:
                    missing.append(f"{wb_name}!{sheet}.{f}")
        if missing:
            fails += 1
            print(f"FAIL `{name}`: {'; '.join(missing)}")
        else:
            print(f"pass `{name}`")

    for wb in wb_cache.values():
        wb.close()
    if fails:
        print(f"\n{fails} 个模式字段缺失 → 降级 candidate 并登记 profile_conflict")
        return 1
    print("\n全部模式字段校验通过 → verified 直接采用")
    return 0


if __name__ == "__main__":
    sys.exit(main())
